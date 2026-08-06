"""Core logic for cleaning OCR/scanned Word documents with broken text boxes."""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}

EMU_PER_INCH = 914400
ROTATION_TOLERANCE = 450000  # 7.5 degrees in OOXML units (60000 = 1 degree)

LIST_ITEM_RE = re.compile(r"^[（(][一二三四五六七八九十百千0-9]+[）)]")
HEADING_RE = re.compile(r"^《.+》第.+条")
PAGE_NUMBER_RE = re.compile(r"^\d{1,3}$")
SPACED_CHARS_RE = re.compile(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])")


def _tag(ns_prefix: str, local: str) -> str:
    return f"{{{NS[ns_prefix]}}}{local}"


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _text_from_element(element: ET.Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        if _local(node.tag) == "t" and node.text:
            parts.append(node.text)
        if _local(node.tag) == "tab":
            parts.append("\t")
        if _local(node.tag) == "br":
            parts.append("\n")
    return "".join(parts)


def _normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = SPACED_CHARS_RE.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _rotation_from_element(element: ET.Element) -> int:
    for node in element.iter():
        if _local(node.tag) == "xfrm":
            rot = node.get("rot")
            if rot:
                try:
                    return int(rot)
                except ValueError:
                    return 0
    return 0


def _offset_from_anchor(anchor: ET.Element) -> tuple[int, int]:
    x = y = 0
    pos_h = anchor.find("wp:positionH", NS)
    if pos_h is not None:
        offset = pos_h.find("wp:posOffset", NS)
        if offset is not None and offset.text:
            x = int(offset.text)
    pos_v = anchor.find("wp:positionV", NS)
    if pos_v is not None:
        offset = pos_v.find("wp:posOffset", NS)
        if offset is not None and offset.text:
            y = int(offset.text)
    return x, y


def _extent_from_anchor(anchor: ET.Element) -> tuple[int, int]:
    extent = anchor.find("wp:extent", NS)
    if extent is None:
        return 0, 0
    cx = extent.get("cx")
    cy = extent.get("cy")
    return (int(cx) if cx else 0, int(cy) if cy else 0)


def _is_rotated(rotation: int) -> bool:
    rotation = rotation % 21600000
    return (
        abs(rotation - 5400000) <= ROTATION_TOLERANCE
        or abs(rotation - 16200000) <= ROTATION_TOLERANCE
    )


def _unrotate_text(text: str, rotation: int) -> str:
    if not _is_rotated(rotation):
        return text

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return text

    rotation = rotation % 21600000
    if abs(rotation - 5400000) <= ROTATION_TOLERANCE:
        # 90° clockwise in Word: read columns right-to-left
        max_len = max(len(line) for line in lines)
        chars: list[str] = []
        for col in range(max_len):
            for row in range(len(lines) - 1, -1, -1):
                if col < len(lines[row]):
                    chars.append(lines[row][col])
        return "".join(chars)

    if abs(rotation - 16200000) <= ROTATION_TOLERANCE:
        # 270° / 90° counter-clockwise
        max_len = max(len(line) for line in lines)
        chars = []
        for col in range(max_len - 1, -1, -1):
            for row in range(len(lines)):
                if col < len(lines[row]):
                    chars.append(lines[row][col])
        return "".join(chars)

    return text


@dataclass(order=True)
class TextBlock:
    sort_y: int
    sort_x: int
    text: str
    source: str = "unknown"
    rotation: int = 0
    sequence: int = 0
    is_page_number: bool = False
    is_heading: bool = False
    is_list_item: bool = False


@dataclass
class CleanStats:
    input_blocks: int = 0
    output_paragraphs: int = 0
    removed_page_numbers: int = 0
    unrotated_blocks: int = 0
    textbox_blocks: int = 0
    warnings: list[str] = field(default_factory=list)


class DocxCleaner:
    def __init__(self, keep_page_numbers: bool = False):
        self.keep_page_numbers = keep_page_numbers

    def clean_file(self, input_path: Path, output_path: Path) -> CleanStats:
        input_path = Path(input_path)
        output_path = Path(output_path)
        blocks = self._extract_blocks(input_path)
        paragraphs = self._merge_blocks(blocks)
        self._write_document(paragraphs, output_path)

        stats = CleanStats(
            input_blocks=len(blocks),
            output_paragraphs=len(paragraphs),
            removed_page_numbers=sum(1 for b in blocks if b.is_page_number),
            unrotated_blocks=sum(1 for b in blocks if _is_rotated(b.rotation)),
            textbox_blocks=sum(1 for b in blocks if b.source == "textbox"),
        )
        return stats

    def _extract_blocks(self, input_path: Path) -> list[TextBlock]:
        blocks: list[TextBlock] = []
        sequence = 0

        with zipfile.ZipFile(input_path) as zf:
            xml_names = [
                name
                for name in zf.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            ]
            for xml_name in sorted(xml_names):
                root = ET.fromstring(zf.read(xml_name))
                blocks.extend(self._extract_from_root(root, xml_name, sequence))
                sequence += 10000

        blocks.extend(self._extract_python_docx_fallback(input_path, sequence))
        blocks = self._deduplicate_blocks(blocks)
        blocks.sort()
        return blocks

    def _extract_from_root(
        self, root: ET.Element, xml_name: str, base_sequence: int
    ) -> list[TextBlock]:
        blocks: list[TextBlock] = []
        sequence = base_sequence

        for anchor in root.findall(".//wp:anchor", NS):
            rotation = _rotation_from_element(anchor)
            x, y = _offset_from_anchor(anchor)
            for txbx_content in anchor.findall(".//w:txbxContent", NS):
                text = _normalize_text(_text_from_element(txbx_content))
                if not text:
                    continue
                text = _unrotate_text(text, rotation)
                block = self._make_block(
                    text=text,
                    x=x,
                    y=y,
                    rotation=rotation,
                    source="textbox",
                    sequence=sequence,
                )
                if block:
                    blocks.append(block)
                    sequence += 1

        # Legacy VML text boxes (common in WPS / older converters)
        for vml_textbox in root.findall(".//v:textbox", NS):
            for txbx_content in vml_textbox.findall(".//w:txbxContent", NS):
                text = _normalize_text(_text_from_element(txbx_content))
                if not text:
                    continue
                block = self._make_block(
                    text=text,
                    x=sequence,
                    y=sequence,
                    rotation=0,
                    source="textbox-vml",
                    sequence=sequence,
                )
                if block:
                    blocks.append(block)
                    sequence += 1

        for drawing in root.findall(".//w:drawing", NS):
            parent_anchor = None
            for ancestor in root.iter():
                if drawing in list(ancestor):
                    if _local(ancestor.tag) == "anchor":
                        parent_anchor = ancestor
                        break
            if parent_anchor is not None:
                continue

            inline = drawing.find("wp:inline", NS)
            if inline is None:
                continue
            rotation = _rotation_from_element(inline)
            extent = inline.find("wp:extent", NS)
            x = y = 0
            if extent is not None:
                x = int(extent.get("cx") or 0)
                y = int(extent.get("cy") or 0)
            for txbx_content in drawing.findall(".//w:txbxContent", NS):
                text = _normalize_text(_text_from_element(txbx_content))
                if not text:
                    continue
                text = _unrotate_text(text, rotation)
                block = self._make_block(
                    text=text,
                    x=x,
                    y=y,
                    rotation=rotation,
                    source="textbox-inline",
                    sequence=sequence,
                )
                if block:
                    blocks.append(block)
                    sequence += 1

        for para in root.findall(".//w:body/w:p", NS):
            text = _normalize_text(_text_from_element(para))
            if not text:
                continue
            if self._paragraph_has_floating_shape(para):
                continue
            block = self._make_block(
                text=text,
                x=0,
                y=sequence,
                rotation=0,
                source="paragraph",
                sequence=sequence,
            )
            if block:
                blocks.append(block)
                sequence += 1

        for table in root.findall(".//w:tbl", NS):
            for row in table.findall(".//w:tr", NS):
                row_texts: list[str] = []
                for cell in row.findall(".//w:tc", NS):
                    cell_text = _normalize_text(_text_from_element(cell))
                    if cell_text:
                        row_texts.append(cell_text)
                if row_texts:
                    text = " | ".join(row_texts)
                    block = self._make_block(
                        text=text,
                        x=0,
                        y=sequence,
                        rotation=0,
                        source="table",
                        sequence=sequence,
                    )
                    if block:
                        blocks.append(block)
                        sequence += 1

        return blocks

    def _extract_python_docx_fallback(
        self, input_path: Path, base_sequence: int
    ) -> list[TextBlock]:
        blocks: list[TextBlock] = []
        try:
            doc = Document(input_path)
        except Exception:
            return blocks

        sequence = base_sequence
        for para in doc.paragraphs:
            text = _normalize_text(para.text)
            if text:
                block = self._make_block(
                    text=text,
                    x=0,
                    y=sequence,
                    rotation=0,
                    source="paragraph-fallback",
                    sequence=sequence,
                )
                if block:
                    blocks.append(block)
                    sequence += 1
        return blocks

    def _paragraph_has_floating_shape(self, para: ET.Element) -> bool:
        for node in para.iter():
            if _local(node.tag) in {"drawing", "pict", "txbxContent"}:
                return True
        return False

    def _make_block(
        self,
        text: str,
        x: int,
        y: int,
        rotation: int,
        source: str,
        sequence: int,
    ) -> TextBlock | None:
        if not text:
            return None

        is_page_number = bool(PAGE_NUMBER_RE.match(text))
        if is_page_number and not self.keep_page_numbers:
            return None

        is_heading = bool(HEADING_RE.match(text))
        is_list_item = bool(LIST_ITEM_RE.match(text))

        sort_y = y if y else sequence
        sort_x = x

        return TextBlock(
            sort_y=sort_y,
            sort_x=sort_x,
            text=text,
            source=source,
            rotation=rotation,
            sequence=sequence,
            is_page_number=is_page_number,
            is_heading=is_heading,
            is_list_item=is_list_item,
        )

    def _deduplicate_blocks(self, blocks: Iterable[TextBlock]) -> list[TextBlock]:
        seen: set[str] = set()
        unique: list[TextBlock] = []
        for block in blocks:
            key = re.sub(r"\s+", "", block.text)
            if key in seen:
                continue
            seen.add(key)
            unique.append(block)
        return unique

    def _merge_blocks(self, blocks: list[TextBlock]) -> list[str]:
        if not blocks:
            return []

        paragraphs: list[str] = []
        buffer = blocks[0].text

        for prev, curr in zip(blocks, blocks[1:]):
            if self._should_merge(prev, curr):
                if prev.is_list_item or curr.is_list_item:
                    paragraphs.append(buffer)
                    buffer = curr.text
                elif prev.is_heading or curr.is_heading:
                    paragraphs.append(buffer)
                    buffer = curr.text
                else:
                    joiner = "" if self._join_without_space(prev.text, curr.text) else ""
                    buffer = f"{buffer}{joiner}{curr.text}"
            else:
                paragraphs.append(buffer)
                buffer = curr.text

        paragraphs.append(buffer)
        return [_normalize_text(p) for p in paragraphs if _normalize_text(p)]

    def _should_merge(self, prev: TextBlock, curr: TextBlock) -> bool:
        if curr.is_heading or curr.is_list_item:
            return False
        if prev.is_heading:
            return False
        if prev.is_list_item and curr.is_list_item:
            return False
        if prev.source.startswith("textbox") or curr.source.startswith("textbox"):
            if abs(prev.sort_y - curr.sort_y) <= EMU_PER_INCH // 2:
                return True
            return False
        if prev.source == "paragraph" and curr.source == "paragraph":
            return True
        return False

    @staticmethod
    def _join_without_space(left: str, right: str) -> bool:
        if not left or not right:
            return False
        return ord(left[-1]) > 127 or ord(right[0]) > 127

    def _write_document(self, paragraphs: list[str], output_path: Path) -> None:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(12)

        for text in paragraphs:
            para = doc.add_paragraph()
            if HEADING_RE.match(text):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(text)
                run.bold = True
                run.font.size = Pt(14)
            elif LIST_ITEM_RE.match(text):
                para.paragraph_format.left_indent = Pt(24)
                para.paragraph_format.first_line_indent = Pt(-24)
                para.add_run(text)
            else:
                para.paragraph_format.first_line_indent = Pt(24)
                para.add_run(text)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)


def clean_docx(
    input_path: str | Path,
    output_path: str | Path | None = None,
    keep_page_numbers: bool = False,
) -> CleanStats:
    input_path = Path(input_path)
    if output_path is None:
        output_path = input_path.with_name(f"{input_path.stem}_已整理{input_path.suffix}")
    else:
        output_path = Path(output_path)

    cleaner = DocxCleaner(keep_page_numbers=keep_page_numbers)
    return cleaner.clean_file(input_path, output_path)


def clean_docx_bytes(data: bytes, keep_page_numbers: bool = False) -> tuple[bytes, CleanStats]:
    with zipfile.ZipFile(BytesIO(data)) as zf:
        if "word/document.xml" not in zf.namelist():
            raise ValueError("不是有效的 .docx 文件")

    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as src:
        src.write(data)
        src_path = Path(src.name)

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as dst:
        dst_path = Path(dst.name)

    try:
        stats = clean_docx(src_path, dst_path, keep_page_numbers=keep_page_numbers)
        output_bytes = dst_path.read_bytes()
        return output_bytes, stats
    finally:
        src_path.unlink(missing_ok=True)
        dst_path.unlink(missing_ok=True)
