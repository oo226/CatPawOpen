"""Compare cleaned DOCX text against a reference PDF."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

import pdfplumber
from docx import Document

PUNCT_RE = re.compile(
    r"[\s\u00a0，。、；：""''（）()【】《》<>·…—.,;:!?\-[\]]"
)
SEGMENT_SPLIT_RE = re.compile(r"[。；\n]+")
MIN_SEGMENT_LEN = 6
MAX_MISSING_DISPLAY = 20


@dataclass
class PdfCheckResult:
    coverage_percent: float
    char_coverage_percent: float
    pdf_char_count: int
    docx_char_count: int
    matched_segments: int
    total_segments: int
    missing_segments: list[str] = field(default_factory=list)
    status: str = "ok"  # ok | warning | error
    message: str = ""

    def to_dict(self) -> dict:
        return {
            "coverage_percent": round(self.coverage_percent, 1),
            "char_coverage_percent": round(self.char_coverage_percent, 1),
            "pdf_char_count": self.pdf_char_count,
            "docx_char_count": self.docx_char_count,
            "matched_segments": self.matched_segments,
            "total_segments": self.total_segments,
            "missing_segments": self.missing_segments,
            "status": self.status,
            "message": self.message,
        }


def normalize_for_compare(text: str) -> str:
    return PUNCT_RE.sub("", text)


def extract_pdf_text(data: bytes) -> str:
    parts: list[str] = []
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
    return "\n".join(parts)


def extract_docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_parts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_parts:
                parts.append(" ".join(row_parts))
    return "\n".join(parts)


def split_segments(text: str) -> list[str]:
    segments: list[str] = []
    for part in SEGMENT_SPLIT_RE.split(text):
        part = part.strip()
        if len(normalize_for_compare(part)) >= MIN_SEGMENT_LEN:
            segments.append(part)
    return segments


def _fuzzy_contains(haystack: str, needle: str, threshold: float = 0.85) -> bool:
    if not needle:
        return True
    if needle in haystack:
        return True

    needle_len = len(needle)
    if needle_len < MIN_SEGMENT_LEN:
        return needle in haystack

    window = max(needle_len, MIN_SEGMENT_LEN)
    best = 0
    for i in range(0, max(1, len(haystack) - window + 1)):
        chunk = haystack[i : i + window]
        matches = sum(1 for a, b in zip(chunk, needle) if a == b)
        best = max(best, matches / needle_len)
        if best >= threshold:
            return True
    return False


def compare_with_pdf(pdf_data: bytes, docx_data: bytes) -> PdfCheckResult:
    pdf_text = extract_pdf_text(pdf_data)
    docx_text = extract_docx_text(docx_data)

    pdf_norm = normalize_for_compare(pdf_text)
    docx_norm = normalize_for_compare(docx_text)

    if not pdf_norm:
        return PdfCheckResult(
            coverage_percent=0,
            char_coverage_percent=0,
            pdf_char_count=0,
            docx_char_count=len(docx_norm),
            matched_segments=0,
            total_segments=0,
            status="error",
            message="PDF 中未能提取到文字（可能是纯图片扫描件，需 OCR 后的 PDF）",
        )

    segments = split_segments(pdf_text)
    missing: list[str] = []
    matched = 0

    for segment in segments:
        seg_norm = normalize_for_compare(segment)
        if seg_norm in docx_norm or _fuzzy_contains(docx_norm, seg_norm):
            matched += 1
        else:
            preview = segment if len(segment) <= 80 else segment[:80] + "…"
            missing.append(preview)

    total = len(segments) or 1
    coverage = (matched / total) * 100

    pdf_chars = len(pdf_norm)
    docx_chars = len(docx_norm)
    common = sum(1 for c in set(pdf_norm) if c in docx_norm)
    char_coverage = (common / max(len(set(pdf_norm)), 1)) * 100

    if coverage >= 95:
        status = "ok"
        message = "对照良好，整理结果与 PDF 内容基本一致"
    elif coverage >= 80:
        status = "warning"
        message = "有部分内容可能缺失，请对照下方缺失片段检查"
    else:
        status = "error"
        message = "缺失较多，建议检查原文件或重新 OCR"

    return PdfCheckResult(
        coverage_percent=coverage,
        char_coverage_percent=char_coverage,
        pdf_char_count=pdf_chars,
        docx_char_count=docx_chars,
        matched_segments=matched,
        total_segments=len(segments),
        missing_segments=missing[:MAX_MISSING_DISPLAY],
        status=status,
        message=message,
    )
