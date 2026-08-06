"""Generate a synthetic messy docx and verify the cleaner."""

from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from cleaner import clean_docx

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": NS_W}


def _make_minimal_docx(paragraphs: list[str], extra_xml: str = "") -> bytes:
    body_parts = []
    for text in paragraphs:
        body_parts.append(
            f'<w:p xmlns:w="{NS_W}"><w:r><w:t>{text}</w:t></w:r></w:p>'
        )
    body_parts.append(extra_xml)
    body = "".join(body_parts)

    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{NS_W}"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
  <w:body>
    {body}
    <w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>
  </w:body>
</w:document>"""

    content_types = """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

    rels = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

    doc_rels = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>"""

    from io import BytesIO

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/_rels/document.xml.rels", doc_rels)
        zf.writestr("word/document.xml", document_xml)
    return buf.getvalue()


def test_basic_paragraphs(tmp_dir: Path) -> None:
    data = _make_minimal_docx(
        [
            "下列事项由业主共同决定：",
            "（一）制定和修改业主大会议事规则；",
            "（二）制定和修改管理规约；",
        ]
    )
    input_path = tmp_dir / "sample.docx"
    output_path = tmp_dir / "sample_已整理.docx"
    input_path.write_bytes(data)

    stats = clean_docx(input_path, output_path)
    assert stats.output_paragraphs >= 3

    from docx import Document

    doc = Document(output_path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("业主共同决定" in t for t in texts)
    assert any("（一）" in t for t in texts)
    print("test_basic_paragraphs: OK")


def test_spaced_chinese(tmp_dir: Path) -> None:
    data = _make_minimal_docx(["选 举 业 主 委 员 会"])
    input_path = tmp_dir / "spaced.docx"
    output_path = tmp_dir / "spaced_已整理.docx"
    input_path.write_bytes(data)

    clean_docx(input_path, output_path)
    from docx import Document

    doc = Document(output_path)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "选举业主委员会" in text.replace(" ", "")
    print("test_spaced_chinese: OK")


def main() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        test_basic_paragraphs(tmp_dir)
        test_spaced_chinese(tmp_dir)
        print("All tests passed.")


if __name__ == "__main__":
    main()
